"""导出明细表与汇总表到 Excel。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.aggregator import DetailRow, SummaryRow


def export_excel(
    path: str | Path,
    detail_rows: list[DetailRow],
    summary_rows: list[SummaryRow],
) -> Path:
    """写入 Excel：汇总表在前，明细表在后。"""
    out = Path(path)
    wb = Workbook()

    # --- 汇总 ---
    ws_sum = wb.active
    ws_sum.title = "汇总"
    headers_sum = ["蔬菜名称", "单位", "斤数"]
    _write_header(ws_sum, headers_sum)
    for i, row in enumerate(summary_rows, start=2):
        ws_sum.cell(i, 1, row.name)
        ws_sum.cell(i, 2, row.unit)
        ws_sum.cell(i, 3, row.jin)
    _autosize(ws_sum, 3)

    # --- 明细 ---
    ws_det = wb.create_sheet("明细")
    headers_det = [
        "来源图片",
        "蔬菜名称",
        "原始数量",
        "原始单位",
        "折合斤数",
        "备注",
    ]
    _write_header(ws_det, headers_det)
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    for i, row in enumerate(detail_rows, start=2):
        ws_det.cell(i, 1, row.source_image)
        ws_det.cell(i, 2, row.name)
        ws_det.cell(i, 3, row.quantity if row.quantity is not None else "")
        ws_det.cell(i, 4, row.unit_raw or "")
        ws_det.cell(i, 5, row.jin if row.jin is not None else "")
        cell_remark = ws_det.cell(i, 6, row.remark)
        if row.remark:
            cell_remark.fill = warn_fill
    _autosize(ws_det, 6)

    wb.save(out)
    return out


def _write_header(ws, headers: list[str]) -> None:
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9EAD3")
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(1, col, title)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        max_len = 8
        for cell in ws[get_column_letter(col)]:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(val) + 2, 40))
        ws.column_dimensions[get_column_letter(col)].width = max_len + 2
