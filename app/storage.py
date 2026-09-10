"""内置真实 Excel 账本（.xlsx）：明细 + 汇总，持久保存。"""

from __future__ import annotations

import os
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import APP_NAME
from app.aggregator import DetailRow, SummaryRow, aggregate
from app.excel_export import export_excel

_LEGACY_MIGRATED = False


def app_root() -> Path:
    """程序安装/运行目录（exe 所在目录；开发时为项目根）。"""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent.parent


def user_data_dir() -> Path:
    """稳定可写的用户数据根目录。

    打包后固定使用 %LOCALAPPDATA%\\蔬菜汇总，避免从微信下载目录直接运行时
    把账本写到临时/只读位置导致 Permission denied。
    开发模式仍用项目根目录，方便本地调试。
    """
    if not getattr(sys, "frozen", False):
        return app_root()
    local = os.environ.get("LOCALAPPDATA")
    if local:
        base = Path(local)
    else:
        base = Path.home() / "AppData" / "Local"
    path = base / APP_NAME
    path.mkdir(parents=True, exist_ok=True)
    return path


def data_dir() -> Path:
    """账本、OCR 调试、自定义菜名等持久文件目录。"""
    path = user_data_dir() / "data"
    path.mkdir(parents=True, exist_ok=True)
    _migrate_legacy_data_once(path)
    return path


def _migrate_legacy_data_once(dest: Path) -> None:
    """若旧版把 data/ 放在 exe 旁，且与新目录不同，则一次性拷贝缺失文件。"""
    global _LEGACY_MIGRATED
    if _LEGACY_MIGRATED or not getattr(sys, "frozen", False):
        return
    _LEGACY_MIGRATED = True
    legacy = app_root() / "data"
    try:
        if not legacy.is_dir():
            return
        if legacy.resolve() == dest.resolve():
            return
        for src in legacy.iterdir():
            if not src.is_file():
                continue
            target = dest / src.name
            if target.exists():
                continue
            try:
                shutil.copy2(src, target)
            except OSError:
                pass
    except OSError:
        pass


def default_ledger_path() -> Path:
    return data_dir() / "蔬菜账本.xlsx"


DETAIL_HEADERS = [
    "编号",
    "入库时间",
    "来源图片",
    "蔬菜名称",
    "原始数量",
    "原始单位",
    "折合斤数",
    "备注",
]

SUMMARY_HEADERS = ["蔬菜名称", "单位", "斤数"]


class Ledger:
    """真实 .xlsx 账本：含「明细」「汇总」两个工作表。"""

    def __init__(self, path: Path | None = None) -> None:
        self.path = path or default_ledger_path()
        self._ensure_file()

    def _ensure_file(self) -> None:
        if self.path.exists():
            self._ensure_sheets()
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "明细"
        self._write_header(ws, DETAIL_HEADERS)
        ws_sum = wb.create_sheet("汇总")
        self._write_header(ws_sum, SUMMARY_HEADERS)
        wb.save(self.path)

    def _ensure_sheets(self) -> None:
        wb = self._load()
        changed = False
        if "明细" not in wb.sheetnames:
            ws = wb.create_sheet("明细", 0)
            self._write_header(ws, DETAIL_HEADERS)
            changed = True
        if "汇总" not in wb.sheetnames:
            ws = wb.create_sheet("汇总")
            self._write_header(ws, SUMMARY_HEADERS)
            changed = True
        if changed:
            wb.save(self.path)

    @staticmethod
    def _write_header(ws, headers: list[str]) -> None:
        bold = Font(bold=True)
        fill = PatternFill("solid", fgColor="D9EAD3")
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(1, col, title)
            cell.font = bold
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

    def _load(self):
        return load_workbook(self.path)

    def _next_id(self, ws) -> int:
        max_id = 0
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0] is None:
                continue
            try:
                max_id = max(max_id, int(row[0]))
            except (TypeError, ValueError):
                continue
        return max_id + 1

    def refresh_summary_sheet(self) -> None:
        """根据明细重算并写回「汇总」工作表。"""
        summary = self.summary()
        wb = self._load()
        if "汇总" in wb.sheetnames:
            del wb["汇总"]
        ws = wb.create_sheet("汇总", 0)
        self._write_header(ws, SUMMARY_HEADERS)
        for i, row in enumerate(summary, start=2):
            ws.cell(i, 1, row.name)
            ws.cell(i, 2, row.unit)
            ws.cell(i, 3, row.jin)
        self._autosize(ws, 3)
        wb.save(self.path)

    def list_source_images(self) -> set[str]:
        """已入库的来源图片文件名。"""
        return {name for name, _, _ in self.list_source_image_stats()}

    def list_source_image_stats(self) -> list[tuple[str, str, int]]:
        """已入库图片统计：(文件名, 最近入库时间, 明细条数)，按时间倒序。"""
        stats: dict[str, list] = {}
        for rec in self.load_records():
            src = (rec.get("source") or "").strip()
            if not src or src == "手工调整":
                continue
            t = (rec.get("time") or "").strip()
            if src not in stats:
                stats[src] = [t, 0]
            stats[src][1] += 1
            if t >= stats[src][0]:
                stats[src][0] = t
        rows = [(name, info[0], info[1]) for name, info in stats.items()]
        rows.sort(key=lambda x: x[1], reverse=True)
        return rows

    def delete_by_source(self, source_image: str) -> int:
        """删除某张来源图片的全部明细（用于覆盖重导）。"""
        source_image = source_image.strip()
        if not source_image:
            return 0
        wb = self._load()
        ws = wb["明细"]
        deleted = 0
        for row_idx in range(ws.max_row, 1, -1):
            cell = ws.cell(row_idx, 3).value
            if cell is not None and str(cell).strip() == source_image:
                ws.delete_rows(row_idx, 1)
                deleted += 1
        wb.save(self.path)
        if deleted:
            self.refresh_summary_sheet()
        return deleted

    def append_details(
        self,
        detail_rows: list[DetailRow],
        *,
        replace_sources: set[str] | None = None,
    ) -> int:
        """追加明细。replace_sources 中的来源图会先删旧再写，避免重复累计。"""
        if not detail_rows:
            return 0
        if replace_sources:
            for src in replace_sources:
                self.delete_by_source(src)
        wb = self._load()
        ws = wb["明细"]
        next_id = self._next_id(ws)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        added = 0
        for row in detail_rows:
            name = (row.name or "").strip()
            if not name:
                continue
            ws.append(
                [
                    next_id,
                    now,
                    row.source_image,
                    name,
                    row.quantity if row.quantity is not None else "",
                    row.unit_raw or "",
                    row.jin if row.jin is not None else "",
                    row.remark or "",
                ]
            )
            next_id += 1
            added += 1
        self._autosize(ws, len(DETAIL_HEADERS))
        wb.save(self.path)
        self.refresh_summary_sheet()
        return added

    def adjust_jin(self, name: str, delta_jin: float, remark: str = "") -> bool:
        """手工加减斤数。delta_jin 为负数表示减斤。"""
        name = (name or "").strip()
        if not name or abs(delta_jin) < 1e-9:
            return False
        from app.parser import clean_vegetable_name

        cleaned = clean_vegetable_name(name) or name
        kind = "减斤" if delta_jin < 0 else "加斤"
        note = remark.strip() or kind
        if kind not in note:
            note = f"{kind}；{note}" if note else kind
        row = DetailRow(
            source_image="手工调整",
            name=cleaned,
            quantity=delta_jin,
            unit_raw="斤",
            jin=delta_jin,
            remark=note,
            include_in_summary=True,
        )
        self.append_details([row])
        return True

    def sheet_matrix(self, sheet_name: str) -> tuple[list[str], list[list]]:
        """读取工作表为 (表头, 数据行)，供表格控件显示。"""
        wb = self._load()
        if sheet_name not in wb.sheetnames:
            return [], []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = ["" if h is None else str(h) for h in rows[0]]
        data: list[list] = []
        for row in rows[1:]:
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            data.append(["" if c is None else c for c in row])
        return headers, data

    def save_detail_matrix(self, headers: list[str], data: list[list]) -> None:
        """把表格控件中的明细写回真实 xlsx，并刷新汇总。"""
        wb = self._load()
        if "明细" in wb.sheetnames:
            del wb["明细"]
        ws = wb.create_sheet("明细")
        use_headers = headers if headers else DETAIL_HEADERS
        self._write_header(ws, use_headers)
        for row in data:
            # 跳过全空行
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            # 对齐列数
            cells = list(row) + [""] * max(0, len(use_headers) - len(row))
            ws.append(cells[: len(use_headers)])
        self._autosize(ws, len(use_headers))
        wb.save(self.path)
        self.refresh_summary_sheet()

    def load_details(self) -> list[DetailRow]:
        wb = self._load()
        ws = wb["明细"]
        rows: list[DetailRow] = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if not values or len(values) < 4 or values[3] is None:
                continue
            name = str(values[3]).strip()
            if not name:
                continue
            qty = values[4] if len(values) > 4 else None
            unit = values[5] if len(values) > 5 else None
            jin = values[6] if len(values) > 6 else None
            remark = values[7] if len(values) > 7 else ""
            try:
                qty_f = float(qty) if qty != "" and qty is not None else None
            except (TypeError, ValueError):
                qty_f = None
            try:
                jin_f = float(jin) if jin != "" and jin is not None else None
            except (TypeError, ValueError):
                jin_f = None
            remark_s = str(remark or "")
            include = True
            if jin_f is None:
                include = False
            if "斤数异常" in remark_s or "待人工核对" in remark_s or "粘连拆分" in remark_s:
                include = False
            rows.append(
                DetailRow(
                    source_image=str(values[2] or ""),
                    name=name,
                    quantity=qty_f,
                    unit_raw=str(unit) if unit not in (None, "") else None,
                    jin=jin_f,
                    remark=remark_s,
                    include_in_summary=include,
                )
            )
        return rows

    def load_records(self) -> list[dict]:
        wb = self._load()
        ws = wb["明细"]
        records: list[dict] = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if not values or values[0] is None:
                continue
            try:
                rid = int(values[0])
            except (TypeError, ValueError):
                continue
            records.append(
                {
                    "id": rid,
                    "time": str(values[1] or ""),
                    "source": str(values[2] or ""),
                    "name": str(values[3] or ""),
                    "quantity": values[4] if len(values) > 4 else "",
                    "unit": values[5] if len(values) > 5 else "",
                    "jin": values[6] if len(values) > 6 else "",
                    "remark": str(values[7] or "") if len(values) > 7 else "",
                }
            )
        return records

    def summary(self) -> list[SummaryRow]:
        return aggregate(self.load_details())

    def delete_by_ids(self, ids: set[int]) -> int:
        if not ids:
            return 0
        wb = self._load()
        ws = wb["明细"]
        deleted = 0
        for row_idx in range(ws.max_row, 1, -1):
            cell_id = ws.cell(row_idx, 1).value
            try:
                rid = int(cell_id)
            except (TypeError, ValueError):
                continue
            if rid in ids:
                ws.delete_rows(row_idx, 1)
                deleted += 1
        wb.save(self.path)
        self.refresh_summary_sheet()
        return deleted

    def delete_by_name(self, name: str) -> int:
        name = name.strip()
        if not name:
            return 0
        wb = self._load()
        ws = wb["明细"]
        deleted = 0
        for row_idx in range(ws.max_row, 1, -1):
            cell_name = ws.cell(row_idx, 4).value
            if cell_name is not None and str(cell_name).strip() == name:
                ws.delete_rows(row_idx, 1)
                deleted += 1
        wb.save(self.path)
        self.refresh_summary_sheet()
        return deleted

    def delete_selected_detail_rows(self, row_indices_0based: list[int]) -> int:
        """按表格中的数据行下标删除（不含表头）。"""
        if not row_indices_0based:
            return 0
        wb = self._load()
        ws = wb["明细"]
        # Excel 行号 = 下标 + 2
        excel_rows = sorted({i + 2 for i in row_indices_0based}, reverse=True)
        deleted = 0
        for r in excel_rows:
            if 2 <= r <= ws.max_row:
                ws.delete_rows(r, 1)
                deleted += 1
        wb.save(self.path)
        self.refresh_summary_sheet()
        return deleted

    def cleanup_names(self) -> int:
        """清洗明细：去前缀、拆粘连、剔非菜名、标记异常斤数。返回改动次数。"""
        from app.parser import (
            MAX_JIN_FOR_SUMMARY,
            _NON_VEGETABLE_NAMES,
            clean_vegetable_name,
            split_vegetable_names,
        )

        wb = self._load()
        ws = wb["明细"]
        changed = 0
        # 收集要追加的拆分行、要删除的行
        to_append: list[list] = []
        drop_rows: list[int] = []

        for row_idx in range(2, ws.max_row + 1):
            name_cell = ws.cell(row_idx, 4)
            raw = name_cell.value
            if raw is None:
                continue
            raw_s = str(raw).strip()
            names = split_vegetable_names(raw_s)
            if not names:
                # 尝试轻清洗
                cleaned = clean_vegetable_name(raw_s)
                if not cleaned or cleaned in _NON_VEGETABLE_NAMES:
                    drop_rows.append(row_idx)
                    changed += 1
                    continue
                names = [cleaned]

            jin_cell = ws.cell(row_idx, 7)
            remark_cell = ws.cell(row_idx, 8)
            try:
                jin_val = float(jin_cell.value) if jin_cell.value not in (None, "") else None
            except (TypeError, ValueError):
                jin_val = None

            # 第一名留下，其余拆成新行；多菜粘连时原数量不可靠，清空待核
            if names[0] != raw_s:
                name_cell.value = names[0]
                changed += 1
            if len(names) > 1:
                ws.cell(row_idx, 5).value = ""
                ws.cell(row_idx, 6).value = ""
                ws.cell(row_idx, 7).value = ""
                old_remark = str(remark_cell.value or "")
                if "粘连拆分" not in old_remark:
                    remark_cell.value = (
                        (old_remark + "；" if old_remark else "") + "粘连拆分待核"
                    )
                changed += 1
                for extra in names[1:]:
                    base = [ws.cell(row_idx, c).value for c in range(1, 9)]
                    base[3] = extra
                    base[4] = ""
                    base[5] = ""
                    base[6] = ""
                    base[7] = "粘连拆分待核"
                    to_append.append(base)
                    changed += 1
                # 已处理斤数，跳过后面异常斤数逻辑的重复清空
                continue

            if jin_val is not None and jin_val > MAX_JIN_FOR_SUMMARY:
                old_remark = str(remark_cell.value or "")
                if "斤数异常" not in old_remark:
                    remark_cell.value = (
                        (old_remark + "；" if old_remark else "") + "斤数异常待核"
                    )
                    changed += 1
                # 异常斤数不参与汇总：清空折合斤数，保留原始数量列供人工看
                jin_cell.value = ""
                changed += 1

        for row_idx in reversed(drop_rows):
            ws.delete_rows(row_idx, 1)

        if to_append:
            next_id = self._next_id(ws)
            for row in to_append:
                row[0] = next_id
                next_id += 1
                ws.append(row)

        wb.save(self.path)
        self.refresh_summary_sheet()
        return changed

    def clear_all(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "明细"
        self._write_header(ws, DETAIL_HEADERS)
        ws_sum = wb.create_sheet("汇总", 0)
        self._write_header(ws_sum, SUMMARY_HEADERS)
        wb.save(self.path)

    def export_to(self, path: str | Path) -> Path:
        details = self.load_details()
        summary = aggregate(details)
        return export_excel(path, details, summary)

    @staticmethod
    def _autosize(ws, ncols: int) -> None:
        for col in range(1, ncols + 1):
            max_len = 8
            letter = get_column_letter(col)
            for cell in ws[letter]:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(val) + 2, 40))
            ws.column_dimensions[letter].width = max_len + 2
